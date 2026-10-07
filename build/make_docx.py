"""Word letterhead template: header/footer are the artwork, body is editable."""
import sys, pathlib
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

hdr_png, ftr_png, out = sys.argv[1:4]
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Mm(210), Mm(297)
s.left_margin = s.right_margin = Mm(20)
s.top_margin, s.bottom_margin = Mm(44), Mm(36)
s.header_distance = Mm(0); s.footer_distance = Mm(0)

st = d.styles['Normal']; st.font.name = 'Montserrat'; st.font.size = Pt(9.5)
st.font.color.rgb = RGBColor(0x0F, 0x1E, 0x2A)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Montserrat'); st.element.rPr.rFonts.set(qn('w:cs'), 'Noto Kufi Arabic')
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.3

from lxml import etree
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'

def page_anchored(par, png, h_mm, y_mm):
    """Insert picture as a floating, behind-text image positioned relative to the PAGE (0, y)."""
    par.paragraph_format.space_after = Pt(0); par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.line_spacing = 1.0
    run = par.add_run(); run.add_picture(png, width=Mm(210), height=Mm(h_mm))
    inline = run._r.find('.//{%s}inline' % WP)
    ns = {'wp': WP}
    anchor = etree.SubElement(inline.getparent(), '{%s}anchor' % WP, nsmap={'wp': WP})
    for k, v in dict(distT='0', distB='0', distL='0', distR='0', simplePos='0', relativeHeight='251658240',
                     behindDoc='1', locked='0', layoutInCell='1', allowOverlap='1').items():
        anchor.set(k, v)
    etree.SubElement(anchor, '{%s}simplePos' % WP, x='0', y='0')
    ph = etree.SubElement(anchor, '{%s}positionH' % WP, relativeFrom='page'); etree.SubElement(ph, '{%s}posOffset' % WP).text = '0'
    pv = etree.SubElement(anchor, '{%s}positionV' % WP, relativeFrom='page'); etree.SubElement(pv, '{%s}posOffset' % WP).text = str(int(Mm(y_mm)))
    for tag in ('extent',):
        anchor.append(inline.find('{%s}%s' % (WP, tag)))
    etree.SubElement(anchor, '{%s}effectExtent' % WP, l='0', t='0', r='0', b='0')
    etree.SubElement(anchor, '{%s}wrapNone' % WP)
    for tag in ('docPr', 'cNvGraphicFramePr'):
        anchor.append(inline.find('{%s}%s' % (WP, tag)))
    anchor.append(inline.find('{http://schemas.openxmlformats.org/drawingml/2006/main}graphic'))
    inline.getparent().remove(inline)

page_anchored(s.header.paragraphs[0], hdr_png, 36, 0)
page_anchored(s.footer.paragraphs[0], ftr_png, 26, 297 - 26)

p = d.add_paragraph(); r = p.add_run('Ref: '); r.bold = True; p.add_run('NCL/____/____')
p.add_run('\t\t\t\t\t\t'); r = p.add_run('Date: '); r.bold = True; p.add_run('__ / __ / ____')
p = d.add_paragraph(); r = p.add_run('Recipient name'); r.bold = True
p.add_run('\nCompany name\nCity, Iraq')
d.add_paragraph('Dear Sir / Madam,')
p = d.add_paragraph(); r = p.add_run('Subject: '); r.bold = True; r2 = p.add_run('Write the subject here'); r2.bold = True
d.add_paragraph('Start typing your letter here. The header and footer are part of the page template and repeat on every page.')
d.add_paragraph('Yours faithfully,')
p = d.add_paragraph(); r = p.add_run('Full Name'); r.bold = True; p.add_run('\nJob Title\nNew City Light')
d.core_properties.title = 'New City Light – Letterhead A4'; d.core_properties.author = 'New City Light'
d.save(out); print('saved', out)
