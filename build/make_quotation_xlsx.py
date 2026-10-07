from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XImg
from openpyxl.worksheet.page import PageMargins
from PIL import Image
import io

NAVY, ORANGE, TINT, LINE = '0F1E2A', 'F36B24', 'F1F2F3', 'D9DDE0'
wb = Workbook(); ws = wb.active; ws.title = 'Quotation'
widths = {'A': 6, 'B': 54, 'C': 10, 'D': 10, 'E': 15, 'F': 17}
for k, v in widths.items(): ws.column_dimensions[k].width = v
f = lambda **k: Font(name='Montserrat', **k)
thin = Side(style='thin', color=LINE)

def put(cell, v, font=None, fill=None, al=None, border=None):
    c = ws[cell]; c.value = v
    if font: c.font = font
    if fill: c.fill = PatternFill('solid', fgColor=fill)
    if al: c.alignment = al
    if border: c.border = border
    return c

# logos
def img(path, h_px):
    im = Image.open(path).convert('RGBA'); w = int(im.width * h_px / im.height)
    im = im.resize((w * 3, h_px * 3), Image.LANCZOS); b = io.BytesIO(); im.save(b, 'PNG'); b.seek(0)
    x = XImg(b); x.width, x.height = w, h_px; return x
ws.add_image(img('brand/logos/png/ncl-horizontal-full-colour.png', 40), 'A1')
ar = img('brand/logos/png/ncl-arabic-navy.png', 40); ws.add_image(ar, 'E1')
for r in (1, 2): ws.row_dimensions[r].height = 22
ws.row_dimensions[3].height = 6
for col in 'ABCDEF': ws[f'{col}3'].fill = PatternFill('solid', fgColor=ORANGE)

put('A5', 'QUOTATION', f(size=20, bold=True, color=NAVY)); ws.merge_cells('A5:C5')
put('E5', 'عرض سعر', Font(name='Noto Kufi Arabic', size=18, bold=True, color=ORANGE), al=Alignment(horizontal='right')); ws.merge_cells('E5:F5')
ws.row_dimensions[5].height = 30

lab = f(size=7, bold=True, color='BD470B'); val = f(size=10, color=NAVY)
def field(r, c1, c2, en, ar_, merge=None):
    put(f'{c1}{r}', f'{en}  /  {ar_}', lab)
    put(f'{c1}{r+1}', None, val, border=Border(bottom=thin))
    if merge: ws.merge_cells(f'{c1}{r+1}:{merge}{r+1}')
field(7, 'A', None, 'QUOTATION NO.', 'رقم العرض', 'B'); field(7, 'C', None, 'DATE', 'التاريخ', 'D'); field(7, 'E', None, 'VALID UNTIL', 'صالح لغاية', 'F')
field(9, 'A', None, 'PREPARED BY', 'أعدّه', 'B'); field(9, 'C', None, 'CURRENCY', 'العملة', 'D'); field(9, 'E', None, 'PAYMENT TERMS', 'شروط الدفع', 'F')
field(11, 'A', None, 'QUOTED TO – COMPANY', 'السادة – الشركة', 'B'); field(11, 'C', None, 'ATTENTION', 'عناية', 'D'); field(11, 'E', None, 'PHONE', 'الهاتف', 'F')
field(13, 'A', None, 'PROJECT NAME', 'اسم المشروع', 'B'); field(13, 'C', None, 'LOCATION', 'الموقع', 'D'); field(13, 'E', None, 'REFERENCE', 'المرجع', 'F')

hr = 16
heads = [('No.\nت', 'center'), ('Description\nالوصف', 'left'), ('Unit\nالوحدة', 'center'), ('Qty\nالكمية', 'right'), ('Unit price\nسعر الوحدة', 'right'), ('Total\nالمجموع', 'right')]
for i, (h, a) in enumerate(heads):
    c = ws.cell(row=hr, column=i + 1, value=h)
    c.font = f(size=8, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY)
    c.alignment = Alignment(horizontal=a, vertical='center', wrap_text=True)
ws.row_dimensions[hr].height = 30
N = 18
for i in range(N):
    r = hr + 1 + i
    ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(row=r, column=1).font = f(size=9, color='5E6B76')
    ws.cell(row=r, column=2).alignment = Alignment(vertical='center', wrap_text=True)
    ws.cell(row=r, column=3).alignment = Alignment(horizontal='center', vertical='center')
    for col in (4, 5): ws.cell(row=r, column=col).alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(row=r, column=6, value=f'=IF(OR(D{r}="",E{r}=""),"",D{r}*E{r})').alignment = Alignment(horizontal='right', vertical='center')
    ws.cell(row=r, column=5).number_format = '#,##0.00'; ws.cell(row=r, column=6).number_format = '#,##0.00'
    for col in range(1, 7):
        c = ws.cell(row=r, column=col); c.border = Border(bottom=thin)
        if col > 1: c.font = f(size=9, color=NAVY)
        if i % 2: c.fill = PatternFill('solid', fgColor='FAFAFA')
    ws.row_dimensions[r].height = 22
last = hr + N
t = last + 2
rows = [('Subtotal  /  المجموع الفرعي', f'=SUM(F{hr+1}:F{last})', None),
        ('Discount %  /  نسبة الخصم', None, 0),
        ('Discount  /  الخصم', f'=F{t}*E{t+1}', None),
        ('Tax / VAT %  /  نسبة الضريبة', None, 0),
        ('Tax / VAT  /  الضريبة', f'=(F{t}-F{t+2})*E{t+3}', None)]
for k, (label, formula, pct) in enumerate(rows):
    r = t + k
    ws.merge_cells(f'B{r}:D{r}')
    put(f'B{r}', label, f(size=8, bold=True, color='5E6B76'), al=Alignment(horizontal='right'))
    if pct is not None:
        c = ws[f'E{r}']; c.value = pct; c.number_format = '0.0%'; c.font = f(size=10, color=NAVY); c.fill = PatternFill('solid', fgColor=TINT)
        c.alignment = Alignment(horizontal='right')
    else:
        c = ws[f'F{r}']; c.value = formula; c.number_format = '#,##0.00'; c.font = f(size=10, bold=True, color=NAVY); c.alignment = Alignment(horizontal='right')
g = t + len(rows)
ws.merge_cells(f'B{g}:E{g}')
put(f'B{g}', 'GRAND TOTAL  /  الإجمالي', f(size=10, bold=True, color='FFFFFF'), NAVY, Alignment(horizontal='right', vertical='center'))
put(f'F{g}', f'=F{t}-F{t+2}+F{t+4}', f(size=12, bold=True, color=ORANGE), NAVY, Alignment(horizontal='right', vertical='center'))
ws[f'F{g}'].number_format = '#,##0.00'; ws.row_dimensions[g].height = 26
for col in 'CDE': ws[f'{col}{g}'].fill = PatternFill('solid', fgColor=NAVY)

tr = g + 2
put(f'A{tr}', 'TERMS & CONDITIONS  /  الشروط والأحكام', f(size=8, bold=True, color=NAVY))
terms = ['1. Prices are valid for the period stated above and exclude any taxes or duties not shown.',
         '2. Delivery: ____ weeks from order confirmation and receipt of advance payment.',
         "3. Warranty: as per the manufacturer's terms for each product supplied.",
         '4. Any change to the scope or quantities will be quoted separately.']
for i, tx in enumerate(terms):
    put(f'A{tr+1+i}', tx, f(size=8, color='33434F')); ws.merge_cells(f'A{tr+1+i}:F{tr+1+i}')
s = tr + 6
for col, (en, ar_) in zip('ACE', [('PREPARED BY', 'أعدّه'), ('APPROVED BY', 'اعتمده'), ('COMPANY STAMP', 'ختم الشركة')]):
    put(f'{col}{s}', f'{en}  /  {ar_}', lab)
for col, span in (('A', 'B'), ('C', 'D'), ('E', 'F')):
    ws.merge_cells(f'{col}{s+1}:{span}{s+3}')
    for rr in range(s + 1, s + 4):
        for cc in (col, span):
            ws[f'{cc}{rr}'].border = Border(left=thin if cc == col else None, right=thin if cc == span else None,
                                            top=thin if rr == s + 1 else None, bottom=thin if rr == s + 3 else None)
for rr in range(s + 1, s + 4): ws.row_dimensions[rr].height = 18

b = s + 5
band = f(size=8, color='FFFFFF'); bandb = f(size=8, bold=True, color='FFFFFF')
for rr in range(b, b + 4):
    for col in 'ABCDEF': ws[f'{col}{rr}'].fill = PatternFill('solid', fgColor=NAVY)
for col in 'ABCDEF': ws[f'{col}{b}'].fill = PatternFill('solid', fgColor=ORANGE)
ws.row_dimensions[b].height = 4
put(f'A{b+1}', 'New City Light  ·  Baghdad – Al-Mansour – Al-Dawoodi, Iraq  ·  بغداد – المنصور – الداوودي', bandb); ws.merge_cells(f'A{b+1}:F{b+1}')
put(f'A{b+2}', '+964 783 577 6393  ·  +964 786 555 3648  ·  +964 780 517 7717', band); ws.merge_cells(f'A{b+2}:F{b+2}')
put(f'A{b+3}', 'admin@dulights.com  ·  www.dulights.com  ·  LIGHTING THE WAY', band); ws.merge_cells(f'A{b+3}:F{b+3}')
for rr in (b + 1, b + 2, b + 3): ws.row_dimensions[rr].height = 17

ws.print_area = f'A1:F{b+3}'
ws.page_setup.paperSize = ws.PAPERSIZE_A4; ws.page_setup.orientation = 'portrait'
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = PageMargins(left=.4, right=.4, top=.4, bottom=.4)
ws.sheet_view.showGridLines = False
ws.freeze_panes = None
import os; os.makedirs('deliverables/05-Quotation', exist_ok=True)
wb.save('deliverables/05-Quotation/NCL-Quotation-A4-Excel-template.xlsx'); print('saved')
